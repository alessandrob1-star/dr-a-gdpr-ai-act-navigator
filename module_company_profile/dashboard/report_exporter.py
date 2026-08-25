"""Generate deterministic compliance assessment reports as DOCX or PDF."""

from __future__ import annotations

from datetime import UTC, datetime
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any

from module_company_profile.dashboard.report_localizer import (
    localize_result,
    localize_timeline_item,
    localize_ui,
)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_FILL = "F2F4F7"
OFFICIAL_TIMELINE_URL = "https://digital-strategy.ec.europa.eu/en/policies/regulatory-framework-ai"
SIMPLIFICATION_AGREEMENT_URL = (
    "https://www.consilium.europa.eu/en/press/press-releases/2026/05/07/"
    "artificial-intelligence-council-and-parliament-agree-to-simplify-and-streamline-rules/"
)


def _register_pdf_fonts() -> tuple[str, str]:
    """Register a Unicode-capable font pair available on the host OS.

    DejaVu Sans is installed in the container image; Arial is the native
    fallback on Windows (C:/Windows/Fonts) and macOS
    (/System/Library/Fonts/Supplemental, with the legacy /Library/Fonts
    location kept for older systems). Helvetica remains a last resort for
    unusually minimal environments and does not cover non-Latin scripts.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = (
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ),
        (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
    )
    for regular_path, bold_path in candidates:
        if regular_path.is_file() and bold_path.is_file():
            if "NavigatorSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("NavigatorSans", str(regular_path)))
                pdfmetrics.registerFont(TTFont("NavigatorSans-Bold", str(bold_path)))
            return "NavigatorSans", "NavigatorSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def _report_data(evaluation: dict[str, Any]) -> dict[str, Any]:
    dashboard = evaluation.get("dashboard_context") or {}
    memory = evaluation.get("company_memory") or {}
    return {
        "company": dashboard.get("company_name") or "-",
        "generated": datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC"),
        "score": dashboard.get("compliance_score") or {},
        "tags": dashboard.get("relevance_tags") or [],
        "warnings": dashboard.get("risk_warnings") or [],
        "missing_controls": (memory.get("controls") or {}).get("missing_or_to_verify") or [],
        "timeline": dashboard.get("compliance_timeline") or [],
        "events": dashboard.get("matched_regulatory_events") or [],
    }


def build_docx_report(evaluation: dict[str, Any], language: str = "en") -> bytes:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    data = _report_data(evaluation)

    def tr(key: str) -> str:
        return localize_ui(language, key)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = (
        Inches(1)
    )
    section.header_distance = section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = f"Dr. G.D.P.R. & AI Act navigator | {tr('export_report')}"
    header.style = normal
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(90, 100, 112)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(tr("report_disclaimer")).font.size = Pt(8)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(tr("export_report").upper())
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run(data["company"])
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(55, 55, 55)
    for label, value in (
        (tr("company_profile"), data["company"]),
        (
            tr("compliance_score"),
            f"{data['score'].get('score', '-')}/100 - "
            f"{localize_result(language, data['score'].get('band', '-'))}",
        ),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    generated = doc.add_paragraph(data["generated"])
    generated.runs[0].font.size = Pt(9)
    generated.runs[0].font.color.rgb = RGBColor(90, 100, 112)

    doc.add_heading(tr("dashboard_score_title"), level=1)
    doc.add_paragraph(tr("dashboard_score_text"))
    doc.add_paragraph(tr("report_disclaimer"))
    if data["tags"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{tr('relevant_tags')}: ").bold = True
        p.add_run(", ".join(str(tag).replace("_", " ") for tag in data["tags"]))

    doc.add_heading(tr("warnings_controls"), level=1)
    for warning in data["warnings"]:
        p = doc.add_paragraph()
        p.add_run(
            f"{localize_result(language, warning.get('level', '')).upper()} - "
            f"{localize_result(language, warning.get('title', ''))}"
        ).bold = True
        doc.add_paragraph(localize_result(language, warning.get("message") or ""))
        action = warning.get("recommended_action")
        if action:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{tr('recommended_action')}: ").bold = True
            p.add_run(localize_result(language, action))

    doc.add_heading(tr("missing_controls"), level=1)
    controls = data["missing_controls"] or [tr("no_missing_controls")]
    for control in controls:
        doc.add_paragraph(localize_result(language, control), style="List Bullet")

    doc.add_heading(tr("compliance_timeline"), level=1)
    doc.add_paragraph(tr("timeline_notice"))
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths_dxa = [1656, 2088, 5616]
    widths = [Inches(value / 1440) for value in widths_dxa]
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")
    table_properties.append(table_indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for index, text in enumerate(("", "", tr("compliance_timeline"))):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].runs[0].bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), LIGHT_FILL)
        cell._tc.get_or_add_tcPr().append(shading)
    for item in data["timeline"]:
        cells = table.add_row().cells
        status, milestone_title = localize_timeline_item(language, item)
        values = (item.get("date", ""), status, milestone_title)
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(cell_width)
            cell_width.set(qn("w:w"), str(widths_dxa[index]))
            cell_width.set(qn("w:type"), "dxa")
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
    source_p = doc.add_paragraph()
    source_p.paragraph_format.space_before = Pt(4)
    source_p.paragraph_format.space_after = Pt(4)
    source_p.add_run(f"{tr('source_label')}: {OFFICIAL_TIMELINE_URL}").font.size = Pt(9)
    if any(item.get("basis") == "political_agreement" for item in data["timeline"]):
        agreement_p = doc.add_paragraph()
        agreement_p.paragraph_format.space_after = Pt(4)
        agreement_p.add_run(f"{tr('source_label')}: {SIMPLIFICATION_AGREEMENT_URL}").font.size = Pt(
            9
        )

    doc.add_heading(tr("regulatory_events"), level=1)
    for event in data["events"][:6]:
        p = doc.add_paragraph()
        p.add_run(
            f"{localize_result(language, str(event.get('priority', '')).lower())} "
            f"({event.get('priority_score', '-')}/100): "
        ).bold = True
        p.add_run(str(event.get("title") or ""))
        if event.get("url"):
            doc.add_paragraph(str(event["url"]))

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def build_pdf_report(evaluation: dict[str, Any], language: str = "en") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    data = _report_data(evaluation)

    def tr(key: str) -> str:
        return localize_ui(language, key)

    def safe(value: Any) -> str:
        return escape(str(value or ""))

    regular_font, bold_font = _register_pdf_fonts()
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles["BodyText"].fontName = regular_font
    styles["BodyText"].fontSize = 10.5
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 6
    for name, size, color in (
        ("Heading1", 16, colors.HexColor(f"#{BLUE}")),
        ("Heading2", 13, colors.HexColor(f"#{BLUE}")),
    ):
        styles[name].fontName = bold_font
        styles[name].fontSize = size
        styles[name].textColor = color
        styles[name].spaceBefore = 12
        styles[name].spaceAfter = 6

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular_font, 8)
        canvas.setFillColor(colors.HexColor("#5A6470"))
        canvas.drawRightString(7.5 * inch, 0.55 * inch, str(document.page))
        canvas.restoreState()

    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=0.8 * inch,
        title=f"{tr('export_report')} - {data['company']}",
    )
    story = [
        Paragraph(
            safe(tr("export_report").upper()),
            ParagraphStyle(
                "TitleCustom",
                parent=styles["Title"],
                fontName=bold_font,
                fontSize=22,
                leading=26,
                spaceAfter=4,
            ),
        ),
        Paragraph(
            safe(data["company"]),
            ParagraphStyle(
                "Subtitle",
                parent=styles["BodyText"],
                fontSize=14,
                textColor=colors.HexColor("#373737"),
                spaceAfter=14,
            ),
        ),
        Paragraph(safe(data["generated"]), styles["BodyText"]),
        Paragraph(
            f"<b>{safe(tr('compliance_score'))}:</b> "
            f"{data['score'].get('score', '-')}/100 - "
            f"{safe(localize_result(language, data['score'].get('band', '-')))}",
            styles["BodyText"],
        ),
        Paragraph(safe(tr("dashboard_score_title")), styles["Heading1"]),
        Paragraph(safe(tr("dashboard_score_text")), styles["BodyText"]),
        Paragraph(safe(tr("report_disclaimer")), styles["BodyText"]),
        Paragraph(safe(tr("warnings_controls")), styles["Heading1"]),
    ]
    for warning in data["warnings"]:
        story.append(
            Paragraph(
                f"<b>{safe(localize_result(language, warning.get('level', '')).upper())} - "
                f"{safe(localize_result(language, warning.get('title', '')))}</b>",
                styles["BodyText"],
            )
        )
        story.append(
            Paragraph(
                safe(localize_result(language, warning.get("message") or "")),
                styles["BodyText"],
            )
        )
        if warning.get("recommended_action"):
            story.append(
                Paragraph(
                    f"<b>{safe(tr('recommended_action'))}:</b> "
                    f"{safe(localize_result(language, warning['recommended_action']))}",
                    styles["BodyText"],
                )
            )
    story.append(Paragraph(safe(tr("missing_controls")), styles["Heading1"]))
    controls = data["missing_controls"] or [tr("no_missing_controls")]
    for control in controls:
        story.append(Paragraph(f"- {safe(localize_result(language, control))}", styles["BodyText"]))
    story.append(Paragraph(safe(tr("compliance_timeline")), styles["Heading1"]))
    story.append(Paragraph(safe(tr("timeline_notice")), styles["BodyText"]))
    rows = [["", "", safe(tr("compliance_timeline"))]]
    for item in data["timeline"]:
        status, title = localize_timeline_item(language, item)
        rows.append(
            [
                Paragraph(safe(item.get("date", "")), styles["BodyText"]),
                Paragraph(safe(status), styles["BodyText"]),
                Paragraph(safe(title), styles["BodyText"]),
            ]
        )
    timeline_table = Table(rows, colWidths=[1.0 * inch, 1.45 * inch, 4.05 * inch], repeatRows=1)
    timeline_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_FILL}")),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C8CDD4")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    source_style = ParagraphStyle("Source", parent=styles["BodyText"], fontSize=8.5)
    story.extend(
        [
            timeline_table,
            Spacer(1, 6),
            Paragraph(f"{safe(tr('source_label'))}: {safe(OFFICIAL_TIMELINE_URL)}", source_style),
        ]
    )
    if any(item.get("basis") == "political_agreement" for item in data["timeline"]):
        story.append(
            Paragraph(
                f"{safe(tr('source_label'))}: {safe(SIMPLIFICATION_AGREEMENT_URL)}",
                source_style,
            )
        )
    story.append(Paragraph(safe(tr("regulatory_events")), styles["Heading1"]))
    for event in data["events"][:6]:
        story.append(
            Paragraph(
                f"<b>{safe(localize_result(language, str(event.get('priority', '')).lower()))} "
                f"({safe(event.get('priority_score', '-'))}/100):</b> "
                f"{safe(event.get('title', ''))}",
                styles["BodyText"],
            )
        )
        if event.get("url"):
            story.append(
                Paragraph(
                    safe(event["url"]),
                    ParagraphStyle(
                        "SourceLink",
                        parent=styles["BodyText"],
                        fontSize=8.5,
                        textColor=colors.HexColor(f"#{BLUE}"),
                    ),
                )
            )
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
